from __future__ import annotations

import csv
import html
import io
import json
import os
import re
import textwrap
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Iterable

try:
    import yaml
except ImportError:  # pragma: no cover
    yaml = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

SUPPORTED_ORG_CHART_IMAGE_TYPES = [".png", ".jpg", ".jpeg", ".webp", ".bmp"]

SUPPORTED_COMPANY_FILE_TYPES = [
    ".txt", ".md", ".json", ".yaml", ".yml", ".csv", ".docx", ".pdf", ".xlsx", ".xlsm",
    *SUPPORTED_ORG_CHART_IMAGE_TYPES,
]
SUPPORTED_PERSONAL_FILE_TYPES = [
    ".txt", ".md", ".json", ".yaml", ".yml", ".docx", ".pdf"
]
SUPPORTED_PERSONNEL_MODEL_FILE_TYPES = [
    ".txt", ".md", ".json", ".yaml", ".yml", ".docx", ".pdf", ".zip"
]


def normalize_file_paths(value: object) -> list[str]:
    """Normalize legacy single-path fields and new multi-path fields into a clean list."""
    if value is None:
        return []
    if isinstance(value, Path):
        text = str(value).strip()
        return [text] if text else []
    if isinstance(value, str):
        raw_items = value.replace("\n", ";").split(";")
        return [item.strip() for item in raw_items if item.strip()]
    if isinstance(value, dict):
        return []
    try:
        iterator = iter(value)  # type: ignore[arg-type]
    except TypeError:
        text = str(value).strip()
        return [text] if text else []
    paths: list[str] = []
    for item in iterator:
        for path in normalize_file_paths(item):
            if path and path not in paths:
                paths.append(path)
    return paths


def merge_file_path_values(*values: object) -> list[str]:
    """Merge several path/list values while preserving order and removing duplicates."""
    paths: list[str] = []
    for value in values:
        for path in normalize_file_paths(value):
            if path and path not in paths:
                paths.append(path)
    return paths


def format_file_names(paths: object, empty: str = "未上传") -> str:
    normalized = normalize_file_paths(paths)
    if not normalized:
        return empty
    return "，".join(Path(path).name for path in normalized)

ORG_CHART_KEYWORDS = (
    "组织架构", "组织结构", "组织图", "架构图", "部门架构", "人员架构", "岗位架构",
    "org", "organization", "organisational", "organizational", "structure", "chart", "department",
)

ORG_LEVELS = [
    "董事会/创始层",
    "CEO/总裁层",
    "CXO/高管层",
    "一级部门负责人",
    "二级部门负责人",
    "团队负责人",
    "核心骨干",
    "专业成员",
]

DEFAULT_DEPARTMENTS = [
    "董事会/创始办公室",
    "战略与投资部",
    "人力资源部",
    "财务部",
    "法务与合规部",
    "研发工程部",
    "产品管理部",
    "设计与用户体验部",
    "数据与算法部",
    "信息技术/基础设施部",
    "市场品牌部",
    "销售商务部",
    "客户成功/客服部",
    "运营管理部",
    "供应链/采购部",
    "行政与公共事务部",
]

DEPARTMENT_ROLE_MAP = {
    "董事会/创始办公室": ["董事长", "CEO", "创始人办公室主任", "战略项目经理"],
    "战略与投资部": ["战略分析师", "投资经理", "业务发展经理", "战略规划总监"],
    "人力资源部": ["HRBP", "招聘经理", "组织发展经理", "薪酬绩效经理"],
    "财务部": ["财务经理", "预算分析师", "审计经理", "税务经理"],
    "法务与合规部": ["法务顾问", "合规经理", "隐私专员", "内控经理"],
    "研发工程部": ["后端工程师", "前端工程师", "QA工程师", "架构师", "研发经理"],
    "产品管理部": ["产品经理", "产品总监", "商业产品经理", "项目经理"],
    "设计与用户体验部": ["UX设计师", "UI设计师", "交互设计师", "设计负责人"],
    "数据与算法部": ["数据分析师", "数据科学家", "算法工程师", "数据平台主管"],
    "信息技术/基础设施部": ["运维工程师", "SRE", "网络工程师", "安全工程师"],
    "市场品牌部": ["品牌经理", "内容营销经理", "增长经理", "公关经理"],
    "销售商务部": ["销售经理", "大客户经理", "渠道经理", "商务拓展经理"],
    "客户成功/客服部": ["客服经理", "客户成功经理", "售后支持工程师"],
    "运营管理部": ["运营经理", "项目运营", "流程优化经理", "区域运营负责人"],
    "供应链/采购部": ["采购经理", "供应链经理", "计划专员", "质量经理"],
    "行政与公共事务部": ["行政经理", "公共事务经理", "政府关系经理", "办公室主管"],
}


def slugify(text: str) -> str:
    text = re.sub(r"\s+", "-", (text or "").strip().lower())
    text = re.sub(r"[^\w\-\u4e00-\u9fff]", "", text)
    return text or "company-skill"


def normalize_api_base(model_type: str, api_base: str | None, local_host: str | None = None) -> str | None:
    if model_type != "local":
        return api_base or None
    host = (local_host or api_base or "http://localhost:11434").rstrip("/")
    if host.endswith("/v1"):
        return host
    return f"{host}/v1"



def _json_get(url: str, timeout: float = 1.8) -> dict:
    """Small stdlib-only JSON GET helper for local model discovery."""
    req = urllib.request.Request(url, headers={"Accept": "application/json"}, method="GET")
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8", errors="ignore") or "{}")


def _strip_v1_suffix(host: str) -> str:
    host = (host or "").rstrip("/")
    return host[:-3] if host.endswith("/v1") else host


def _as_openai_v1_base(host: str) -> str:
    host = (host or "").rstrip("/")
    return host if host.endswith("/v1") else f"{host}/v1"


def _append_candidate(candidates: list[tuple[str, str]], service_type: str, host: str) -> None:
    host = (host or "").strip().rstrip("/")
    if not host:
        return
    item = (service_type, host)
    if item not in candidates:
        candidates.append(item)


def detect_local_models(local_host: str | None = None, local_model_type: str | None = None, timeout: float = 1.8) -> list[dict[str, str]]:
    """Detect locally served LLM models from Ollama and OpenAI-compatible endpoints.

    Returns dictionaries containing model_name, model_type, host, api_base, source and endpoint.
    It intentionally avoids external dependencies so the GUI can use it before requirements are installed.
    """
    candidates: list[tuple[str, str]] = []
    current_host = (local_host or "").strip().rstrip("/")
    requested_type = (local_model_type or "").strip().lower()

    if current_host:
        if requested_type == "ollama":
            _append_candidate(candidates, "ollama", _strip_v1_suffix(current_host))
            _append_candidate(candidates, "openai", _as_openai_v1_base(current_host))
        else:
            _append_candidate(candidates, "openai", _as_openai_v1_base(current_host))
            _append_candidate(candidates, "ollama", _strip_v1_suffix(current_host))

    # Common defaults: Ollama, LM Studio, and generic OpenAI-compatible local servers.
    for host in ("http://localhost:11434", "http://127.0.0.1:11434"):
        _append_candidate(candidates, "ollama", host)
        _append_candidate(candidates, "openai", f"{host}/v1")
    for host in ("http://localhost:1234/v1", "http://127.0.0.1:1234/v1", "http://localhost:8000/v1", "http://127.0.0.1:8000/v1"):
        _append_candidate(candidates, "openai", host)

    discovered: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()

    for service_type, host in candidates:
        try:
            if service_type == "ollama":
                base_host = _strip_v1_suffix(host)
                endpoint = f"{base_host}/api/tags"
                body = _json_get(endpoint, timeout=timeout)
                for item in body.get("models", []):
                    name = item.get("name") or item.get("model")
                    if not name:
                        continue
                    key = ("ollama", base_host, str(name))
                    if key in seen:
                        continue
                    seen.add(key)
                    discovered.append({
                        "model_name": str(name),
                        "model_type": "ollama",
                        "host": base_host,
                        "api_base": f"{base_host}/v1",
                        "source": "Ollama /api/tags",
                        "endpoint": endpoint,
                    })
            else:
                api_base = _as_openai_v1_base(host)
                endpoint = f"{api_base}/models"
                body = _json_get(endpoint, timeout=timeout)
                data = body.get("data", [])
                for item in data:
                    name = item.get("id") if isinstance(item, dict) else None
                    if not name:
                        continue
                    lower_host = api_base.lower()
                    if ":1234" in lower_host:
                        model_type = "lm-studio"
                        source = "LM Studio /v1/models"
                    elif ":11434" in lower_host:
                        model_type = "ollama"
                        source = "Ollama OpenAI /v1/models"
                    else:
                        model_type = "其他"
                        source = "OpenAI 兼容 /v1/models"
                    base_host = api_base[:-3] if api_base.endswith("/v1") else api_base
                    key = (model_type, api_base, str(name))
                    if key in seen:
                        continue
                    seen.add(key)
                    discovered.append({
                        "model_name": str(name),
                        "model_type": model_type,
                        "host": base_host,
                        "api_base": api_base,
                        "source": source,
                        "endpoint": endpoint,
                    })
        except Exception:
            # Discovery should be quiet: unavailable ports are normal during scanning.
            continue

    return discovered


def auto_select_local_model(local_host: str | None = None, local_model_type: str | None = None) -> dict[str, str] | None:
    """Return the first detected local model, preferring the user-selected host/type."""
    models = detect_local_models(local_host=local_host, local_model_type=local_model_type)
    return models[0] if models else None


def looks_like_org_chart_file(path: str | Path) -> bool:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    name = file_path.stem.lower()
    if suffix in SUPPORTED_ORG_CHART_IMAGE_TYPES:
        return any(keyword.lower() in name for keyword in ORG_CHART_KEYWORDS) or True
    return any(keyword.lower() in name for keyword in ORG_CHART_KEYWORDS)


def detect_org_chart_files(file_paths: Iterable[str | Path]) -> list[str]:
    return [str(path) for path in file_paths if looks_like_org_chart_file(path)]


def safe_read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "cp936", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _collapse_whitespace(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text or "")).strip()


def _extract_docx_text_with_zip(file_path: Path) -> str:
    """Extract readable text from a .docx file without optional dependencies.

    A .docx file is a ZIP package containing XML parts. This fallback reads
    word/document.xml and common table text nodes so DOCX files are never fed
    to the LLM as raw binary PK data when python-docx is unavailable.
    """
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs: list[str] = []
    with zipfile.ZipFile(file_path, "r") as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError("DOCX 包中缺少 word/document.xml")
        xml_bytes = archive.read("word/document.xml")

    root = ET.fromstring(xml_bytes)
    for paragraph in root.findall(".//w:p", namespace):
        pieces = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
        line = "".join(pieces).strip()
        if line:
            paragraphs.append(line)
    return _collapse_whitespace("\n".join(paragraphs))


def extract_docx_text(file_path: Path, max_chars: int = 6000) -> str:
    """Extract DOCX text with python-docx first, then Open XML ZIP fallback."""
    errors: list[str] = []
    if Document is not None:
        try:
            doc = Document(file_path)
            parts: list[str] = []
            parts.extend(p.text.strip() for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = _collapse_whitespace("\n".join(parts))
            if text:
                return text[:max_chars]
            errors.append("python-docx 未提取到正文")
        except Exception as exc:  # pragma: no cover - depends on user files
            errors.append(f"python-docx 解析失败：{exc}")
    else:
        errors.append("未安装 python-docx")

    try:
        text = _extract_docx_text_with_zip(file_path)
        if text:
            return text[:max_chars]
        errors.append("Open XML 兜底解析未提取到正文")
    except Exception as exc:  # pragma: no cover - depends on user files
        errors.append(f"Open XML 兜底解析失败：{exc}")

    return f"[DOCX解析失败] {file_path.name}。" + "；".join(errors) + "。请确认文件未损坏，并运行 pip install -r requirements.txt。"


def _looks_like_binary_zip_text(text: str) -> bool:
    sample = (text or "")[:400]
    return "PK\x03\x04" in sample or "[Content_Types].xml" in sample or "word/document.xml" in sample


def extract_text_from_file(path: str | Path, max_chars: int = 6000) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if not file_path.exists():
        return f"[文件不存在] {file_path}"

    try:
        if suffix in SUPPORTED_ORG_CHART_IMAGE_TYPES:
            note = "可能是组织架构图，将在生成时作为组织架构输入线索。" if looks_like_org_chart_file(file_path) else "图片文件，当前文本摘要阶段不会执行 OCR。"
            return f"[图片文件] {file_path.name}；{note}"[:max_chars]

        if suffix == ".zip":
            with zipfile.ZipFile(file_path, "r") as archive:
                names = archive.namelist()[:80]
            return ("[压缩包文件] " + file_path.name + "\n包含文件：\n" + "\n".join(names))[:max_chars]

        if suffix in {".txt", ".md", ".py", ".log", ".ini", ".cfg"}:
            return safe_read_text(file_path)[:max_chars]

        if suffix == ".json":
            data = json.loads(safe_read_text(file_path))
            return json.dumps(data, ensure_ascii=False, indent=2)[:max_chars]

        if suffix in {".yaml", ".yml"}:
            raw = safe_read_text(file_path)
            if yaml is None:
                return raw[:max_chars]
            data = yaml.safe_load(raw)
            return json.dumps(data, ensure_ascii=False, indent=2)[:max_chars]

        if suffix == ".csv":
            rows: list[list[str]] = []
            reader = csv.reader(io.StringIO(safe_read_text(file_path)))
            for idx, row in enumerate(reader):
                rows.append(row)
                if idx >= 20:
                    break
            return "\n".join([", ".join(map(str, row)) for row in rows])[:max_chars]

        if suffix in {".xlsx", ".xlsm"} and load_workbook is not None:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            parts: list[str] = []
            for sheet_name in workbook.sheetnames[:3]:
                ws = workbook[sheet_name]
                parts.append(f"[工作表] {sheet_name}")
                for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        parts.append(" | ".join(values))
                    if row_idx >= 15:
                        break
            return "\n".join(parts)[:max_chars]

        if suffix == ".docx":
            return extract_docx_text(file_path, max_chars=max_chars)

        if suffix == ".pdf" and PdfReader is not None:
            reader = PdfReader(str(file_path))
            pages: list[str] = []
            collected = 0
            # PDF text extraction can be slow on large or unusual files. Stop as soon
            # as the caller's preview budget is filled instead of always scanning 8 pages.
            for page in reader.pages[:8]:
                page_text = page.extract_text() or ""
                pages.append(page_text)
                collected += len(page_text)
                if collected >= max_chars:
                    break
            return "\n".join(pages)[:max_chars]

        # Never feed binary Office/ZIP files into the prompt as text.
        if suffix in {".docx", ".xlsx", ".xlsm", ".zip"}:
            return f"[解析失败] {file_path.name}: 当前文件类型需要专用解析器，已阻止按普通文本读取二进制内容。"

        text = safe_read_text(file_path)[:max_chars]
        if _looks_like_binary_zip_text(text):
            return f"[解析失败] {file_path.name}: 检测到 ZIP/Office 二进制内容，已阻止写入提示词。"
        return text
    except Exception as exc:  # pragma: no cover
        return f"[解析失败] {file_path.name}: {exc}"




DEPARTMENT_EXTRA_KEYWORDS = {
    "董事会/创始办公室": ["董事会", "创始", "创始人", "CEO", "总裁", "秘书处", "总办", "战略决策", "经营例会"],
    "战略与投资部": ["战略", "投资", "融资", "并购", "商业分析", "行业研究", "竞品", "战略规划", "BD", "business development"],
    "人力资源部": ["人力", "HR", "HRBP", "招聘", "培训", "绩效", "薪酬", "组织发展", "OD", "员工关系", "人才盘点"],
    "财务部": ["财务", "会计", "预算", "报销", "税务", "审计", "收入", "利润", "现金流", "成本", "财报"],
    "法务与合规部": ["法务", "合同", "合规", "风控", "审查", "诉讼", "隐私", "数据合规", "知识产权", "专利", "内控"],
    "研发工程部": ["研发", "开发", "工程", "代码", "架构", "后端", "前端", "QA", "质量保障", "DevOps", "版本", "接口", "算法实现"],
    "产品管理部": ["产品", "需求", "PRD", "原型", "路线图", "用户故事", "功能设计", "项目管理", "迭代", "验收"],
    "设计与用户体验部": ["设计", "UI", "UX", "交互", "视觉", "用户体验", "原型图", "可用性", "设计系统", "Figma"],
    "数据与算法部": ["数据", "算法", "模型", "机器学习", "AI", "报表", "BI", "指标", "数据仓库", "画像", "A/B"],
    "信息技术/基础设施部": ["IT", "信息技术", "基础设施", "运维", "SRE", "网络", "服务器", "云平台", "安全", "权限", "监控"],
    "市场品牌部": ["市场", "品牌", "营销", "传播", "内容", "公关", "活动", "增长", "投放", "SEO", "社媒"],
    "销售商务部": ["销售", "商务", "客户拓展", "大客户", "渠道", "合同谈判", "回款", "商机", "报价", "KA"],
    "客户成功/客服部": ["客服", "客户成功", "售后", "工单", "续费", "客户满意度", "客户培训", "服务支持", "投诉"],
    "运营管理部": ["运营", "流程", "项目运营", "活动运营", "内容运营", "用户运营", "SOP", "效率", "增长运营"],
    "供应链/采购部": ["供应链", "采购", "物流", "仓储", "库存", "供应商", "交付", "计划", "质量", "物料"],
    "行政与公共事务部": ["行政", "办公室", "公共事务", "政府关系", "后勤", "资产", "物业", "会议室", "差旅", "接待"],
}


def normalize_department_candidates(candidates: Iterable[str] | None = None) -> list[str]:
    """Return a stable, de-duplicated department candidate list for classification."""
    result: list[str] = []
    for item in list(candidates or []) + DEFAULT_DEPARTMENTS:
        name = str(item or "").strip()
        if name and name not in result:
            result.append(name)
    return result


def _department_keywords(department: str) -> list[str]:
    keywords: list[str] = []
    for raw in re.split(r"[/／、,，\s]+", department):
        raw = raw.strip()
        if len(raw) >= 2 and raw not in keywords:
            keywords.append(raw)
    for word in DEPARTMENT_EXTRA_KEYWORDS.get(department, []):
        if word and word not in keywords:
            keywords.append(word)
    for role in DEPARTMENT_ROLE_MAP.get(department, []):
        if role and role not in keywords:
            keywords.append(role)
        for part in re.split(r"[/／、,，\s]+", role):
            if len(part) >= 2 and part not in keywords:
                keywords.append(part)
    return keywords


def classify_employee_department_from_text(text: str, candidates: Iterable[str] | None = None, file_name: str = "") -> dict:
    """Classify an employee profile into the most likely department using local, explainable rules."""
    departments = normalize_department_candidates(candidates)
    combined = f"{file_name}\n{text or ''}"
    combined_lower = combined.lower()
    scores: list[dict] = []
    for department in departments:
        score = 0.0
        evidence: list[str] = []
        if department and department.lower() in combined_lower:
            score += 8.0
            evidence.append(f"命中完整部门名：{department}")
        for keyword in _department_keywords(department):
            keyword_lower = keyword.lower()
            if not keyword_lower:
                continue
            count = combined_lower.count(keyword_lower)
            if count <= 0:
                continue
            weight = 4.0 if keyword in DEPARTMENT_ROLE_MAP.get(department, []) else 2.0
            if keyword in DEPARTMENT_EXTRA_KEYWORDS.get(department, []):
                weight = 3.0
            score += min(count, 5) * weight
            if len(evidence) < 6:
                evidence.append(f"{keyword}×{count}")
        scores.append({"department": department, "score": round(score, 2), "evidence": evidence})
    scores.sort(key=lambda item: item["score"], reverse=True)
    best = scores[0] if scores else {"department": "", "score": 0.0, "evidence": []}
    second_score = scores[1]["score"] if len(scores) > 1 else 0.0
    best_score = float(best.get("score") or 0.0)
    if best_score <= 0:
        confidence = 0.0
        department = ""
        evidence = []
        reason = "未在员工资料中命中明显部门、岗位或职能关键词。"
    else:
        margin = max(best_score - float(second_score), 0.0)
        confidence = max(0.25, min(0.98, (best_score + margin) / (best_score + second_score + 8.0)))
        department = str(best["department"])
        evidence = list(best.get("evidence") or [])
        reason = "；".join(evidence[:5]) or "根据岗位/职能关键词匹配。"
    return {
        "mode": "auto",
        "detected_department": department,
        "department": department,
        "confidence": round(confidence, 3),
        "confidence_percent": round(confidence * 100, 1),
        "reason": reason,
        "evidence": evidence,
        "candidate_scores": scores[:8],
    }


def classify_employee_file_department(path: str | Path, candidates: Iterable[str] | None = None, max_chars: int = 6000) -> dict:
    """Extract text from an employee profile file and classify its department locally."""
    file_path = Path(path)
    text = extract_text_from_file(file_path, max_chars=max_chars)
    result = classify_employee_department_from_text(text, candidates=candidates, file_name=file_path.name)
    result["file_path"] = str(file_path)
    result["file_name"] = file_path.name
    result["text_chars"] = len(text or "")
    return result


def classify_employee_files_department(
    paths: object,
    candidates: Iterable[str] | None = None,
    max_chars: int = 12000,
    max_chars_per_file: int = 2500,
    max_files: int = 20,
    progress_callback: object | None = None,
) -> dict:
    """Extract and merge several employee profile files, then classify the combined profile.

    Multi-file personnel detection is meant to be a quick local preview, not a full
    document-ingestion pass. It therefore caps both files and characters so selecting
    many resumes or large PDFs does not leave the GUI apparently stuck on detection.
    """
    file_paths = normalize_file_paths(paths)
    if not file_paths:
        raise ValueError("未提供员工资料文件。")
    selected_paths = file_paths[:max_files]
    skipped_files = max(0, len(file_paths) - len(selected_paths))
    chunks: list[str] = []
    total_chars = 0
    for index, raw_path in enumerate(selected_paths, start=1):
        file_path = Path(raw_path)
        if callable(progress_callback):
            try:
                progress_callback(f"正在解析员工资料 {index}/{len(selected_paths)}：{file_path.name}")
            except Exception:
                pass
        remaining = max(max_chars - total_chars, 0)
        if remaining <= 0:
            break
        preview_chars = min(max_chars_per_file, remaining)
        text = extract_text_from_file(file_path, max_chars=preview_chars)
        total_chars += len(text or "")
        if text.strip():
            chunks.append(f"【员工资料文件：{file_path.name}】\n{text}")
    combined_text = "\n\n".join(chunks)[:max_chars]
    names = "；".join(Path(path).name for path in selected_paths)
    result = classify_employee_department_from_text(combined_text, candidates=candidates, file_name=names)
    result["file_paths"] = file_paths
    result["file_names"] = [Path(path).name for path in file_paths]
    result["file_count"] = len(file_paths)
    result["processed_file_count"] = len(selected_paths)
    result["skipped_file_count"] = skipped_files
    result["text_chars"] = total_chars
    if skipped_files:
        note = f"为避免识别卡死，仅预览前 {len(selected_paths)} 个文件，已跳过 {skipped_files} 个文件。"
        result["reason"] = f"{result.get('reason') or ''}；{note}".strip("；")
    return result


BATCH_MEMBER_FILE_TYPES = sorted(set(SUPPORTED_PERSONAL_FILE_TYPES) | set(SUPPORTED_PERSONNEL_MODEL_FILE_TYPES))
BATCH_MEMBER_SKILL_NAME_HINTS = (
    "skill", "agent", "persona", "prompt", "模型", "蒸馏", "角色", "口径", "画像",
)
BATCH_MEMBER_METADATA_FILENAMES = {"member.json", "member.yaml", "member.yml", "成员信息.json", "成员信息.yaml", "成员信息.yml"}

BATCH_MEMBER_EXCLUDED_DIR_NAMES = {
    ".git", ".svn", ".hg", ".idea", ".vscode", "__pycache__",
    "node_modules", ".venv", "venv", "env", "dist", "build",
    "site-packages", ".mypy_cache", ".ruff_cache",
}
BATCH_MEMBER_MAX_MEMBERS = 80
BATCH_MEMBER_MAX_FILES_PER_MEMBER = 24
BATCH_MEMBER_MAX_SCAN_ENTRIES_PER_MEMBER = 800
BATCH_MEMBER_MAX_SCAN_DEPTH = 2


def _clean_member_folder_name(name: str) -> str:
    cleaned = re.sub(r"^[\s\d._\-、()（）]+", "", name or "").strip()
    cleaned = cleaned.replace("_", " ").strip()
    return cleaned or (name or "未命名成员")


def _looks_like_member_skill_file(path: Path) -> bool:
    lower_name = path.name.lower()
    return any(hint.lower() in lower_name for hint in BATCH_MEMBER_SKILL_NAME_HINTS)


def _load_member_metadata(folder: Path) -> dict:
    for filename in BATCH_MEMBER_METADATA_FILENAMES:
        candidate = folder / filename
        if not candidate.exists() or not candidate.is_file():
            continue
        try:
            if candidate.suffix.lower() == ".json":
                data = json.loads(candidate.read_text(encoding="utf-8"))
            elif yaml is not None:
                data = yaml.safe_load(candidate.read_text(encoding="utf-8"))
            else:
                data = None
        except Exception:
            data = None
        if isinstance(data, dict):
            return data
    return {}


def _extract_member_field_from_text(text: str, field_names: Iterable[str]) -> str:
    names = "|".join(re.escape(name) for name in field_names)
    patterns = [
        rf"(?:^|[\n\r\-•*])\s*(?:{names})\s*[:：]\s*([^\n\r]+)",
        rf"(?:^|[\n\r])#+\s*(?:{names})\s*[:：]\s*([^\n\r]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text or "", flags=re.I)
        if match:
            value = re.sub(r"[`*_#\[\]{}]", "", match.group(1)).strip()
            value = re.split(r"[；;|]", value)[0].strip()
            if value:
                return value[:80]
    return ""


def _preview_member_text(paths: list[str], max_files: int = 4, max_chars_per_file: int = 1800) -> str:
    chunks: list[str] = []
    for raw_path in paths[:max_files]:
        try:
            file_path = Path(raw_path)
            text = extract_text_from_file(file_path, max_chars=max_chars_per_file).strip()
        except Exception:
            text = ""
        if text:
            chunks.append(f"【{Path(raw_path).name}】\n{text}")
    return "\n\n".join(chunks)


def _is_hidden_or_excluded_dir(path: Path) -> bool:
    name = path.name
    if not name:
        return False
    return name.startswith(".") or name in BATCH_MEMBER_EXCLUDED_DIR_NAMES


def _iter_member_supported_files(
    directory: Path,
    recursive: bool = True,
    max_files: int = BATCH_MEMBER_MAX_FILES_PER_MEMBER,
    max_scan_entries: int = BATCH_MEMBER_MAX_SCAN_ENTRIES_PER_MEMBER,
    max_depth: int = BATCH_MEMBER_MAX_SCAN_DEPTH,
) -> tuple[list[Path], dict]:
    """Safely scan one member directory without materializing the full tree.

    Path.glob("**/*") is convenient, but sorting it forces Python to enumerate the
    entire tree before we can stop. Batch import should be defensive because users
    may accidentally select a large workspace, Downloads folder, or synced drive.
    """
    files: list[Path] = []
    scanned_entries = 0
    skipped_dirs = 0
    truncated = False
    root = Path(directory)

    def walk(current: Path, depth: int) -> None:
        nonlocal scanned_entries, skipped_dirs, truncated
        if truncated:
            return
        if depth > max_depth:
            skipped_dirs += 1
            return
        try:
            entries_iter = current.iterdir()
        except Exception:
            skipped_dirs += 1
            return
        for entry in entries_iter:
            if truncated:
                return
            scanned_entries += 1
            if scanned_entries > max_scan_entries:
                truncated = True
                return
            try:
                if entry.is_dir():
                    if recursive and not entry.is_symlink() and not _is_hidden_or_excluded_dir(entry):
                        walk(entry, depth + 1)
                    else:
                        skipped_dirs += 1
                    continue
                if not entry.is_file() or entry.name.startswith("."):
                    continue
            except Exception:
                continue
            if entry.name in BATCH_MEMBER_METADATA_FILENAMES:
                continue
            if entry.suffix.lower() not in BATCH_MEMBER_FILE_TYPES:
                continue
            files.append(entry)
            if len(files) >= max_files:
                truncated = True
                return

    walk(root, 0)
    return files, {
        "truncated": truncated,
        "scanned_entries": scanned_entries,
        "skipped_dirs": skipped_dirs,
        "max_files": max_files,
        "max_scan_entries": max_scan_entries,
        "max_depth": max_depth,
    }


def collect_member_directory_files(folder: str | Path, recursive: bool = True, max_files: int = BATCH_MEMBER_MAX_FILES_PER_MEMBER) -> dict:
    """Collect supported profile/Skill files from one member folder.

    Batch import uses the convention "one member = one sub-folder". Markdown files are
    treated as profile files unless their file name contains Skill/model hints.
    The scanner is intentionally bounded to avoid hanging on huge directories.
    """
    directory = Path(folder)
    files, scan_meta = _iter_member_supported_files(
        directory,
        recursive=recursive,
        max_files=max_files,
    )
    profile_paths: list[str] = []
    skill_paths: list[str] = []
    for path in files:
        target = skill_paths if _looks_like_member_skill_file(path) else profile_paths
        target.append(str(path))
    return {
        "profile_file_paths": profile_paths,
        "skill_file_paths": skill_paths,
        "all_file_paths": [str(path) for path in files],
        **scan_meta,
    }


def build_member_payload_from_directory(
    folder: str | Path,
    candidates: Iterable[str] | None = None,
    default_level: str = "专业成员",
    default_department: str = "产品与解决方案",
    default_member_type: str = "组织成员",
    default_authority_level: str = "advisory",
    recursive: bool = True,
    progress_callback: object | None = None,
) -> dict:
    """Build an organization member payload from a directory containing one person's files."""
    directory = Path(folder)
    if callable(progress_callback):
        try:
            progress_callback(f"正在读取成员资料文件列表：{directory.name}")
        except Exception:
            pass
    metadata = _load_member_metadata(directory)
    collected = collect_member_directory_files(directory, recursive=recursive)
    profile_paths = normalize_file_paths(metadata.get("profile_file_paths") or metadata.get("file_paths") or collected["profile_file_paths"])
    skill_paths = normalize_file_paths(metadata.get("skill_file_paths") or metadata.get("model_file_paths") or collected["skill_file_paths"])
    all_paths = merge_file_path_values(profile_paths, skill_paths, collected.get("all_file_paths"))
    scan_warning = ""
    if collected.get("truncated"):
        scan_warning = (
            f"成员目录扫描已截断：最多读取 {collected.get('max_files')} 个支持文件，"
            f"或最多扫描 {collected.get('max_scan_entries')} 个目录项。"
        )
    if not all_paths:
        raise ValueError(f"成员目录中没有支持的资料文件：{directory}")

    if callable(progress_callback):
        try:
            progress_callback(f"正在预览成员资料：{directory.name}")
        except Exception:
            pass
    preview_text = _preview_member_text(merge_file_path_values(profile_paths, all_paths), max_files=3, max_chars_per_file=1200)
    detected_name = _extract_member_field_from_text(preview_text, ["姓名", "成员名称", "名称", "Name"])
    detected_position = _extract_member_field_from_text(preview_text, ["岗位", "职位", "位置", "角色", "Position", "Role"])
    detected_department_text = _extract_member_field_from_text(preview_text, ["部门", "所属部门", "插入部门", "Department"])
    detected_level = _extract_member_field_from_text(preview_text, ["层级", "插入层级", "Level"])
    detected_authority = _extract_member_field_from_text(preview_text, ["权限", "权限等级", "Authority"])
    role_notes = metadata.get("role_notes") or metadata.get("notes") or _extract_member_field_from_text(preview_text, ["说明", "补充说明", "口径", "备注"])

    if callable(progress_callback):
        try:
            progress_callback(f"正在识别成员所属部门：{directory.name}")
        except Exception:
            pass
    detection = classify_employee_files_department(
        merge_file_path_values(profile_paths, skill_paths),
        candidates=candidates,
        max_chars=5000,
        max_chars_per_file=1200,
        max_files=5,
        progress_callback=progress_callback,
    )
    auto_department = detection.get("department") or ""
    final_department = (
        metadata.get("department")
        or detected_department_text
        or (auto_department if float(detection.get("confidence") or 0.0) > 0 else "")
        or default_department
    )
    assignment = dict(detection)
    assignment.update({
        "mode": "auto-batch",
        "manual_department": str(final_department),
        "final_department": str(final_department),
        "source_directory": str(directory),
        "scan_truncated": bool(collected.get("truncated")),
        "scan_warning": scan_warning,
        "scanned_entries": collected.get("scanned_entries"),
    })
    if not assignment.get("department"):
        assignment["department"] = str(final_department)

    member_name = metadata.get("member_name") or metadata.get("name") or detected_name or _clean_member_folder_name(directory.name)
    position = metadata.get("position") or detected_position or member_name
    level = metadata.get("level") or detected_level or default_level
    authority_level = metadata.get("authority_level") or detected_authority or default_authority_level
    member_type = metadata.get("member_type") or default_member_type
    return {
        "enabled": True,
        "member_name": str(member_name).strip() or _clean_member_folder_name(directory.name),
        "member_type": str(member_type).strip() or default_member_type,
        "level": str(level).strip() or default_level,
        "department": str(final_department).strip() or default_department,
        "position": str(position).strip() or str(member_name).strip() or _clean_member_folder_name(directory.name),
        "role_notes": str(role_notes or "").strip(),
        "has_peer_models": bool(metadata.get("has_peer_models", True)),
        "profile_file_path": profile_paths[0] if profile_paths else "",
        "profile_file_paths": profile_paths,
        "skill_file_path": skill_paths[0] if skill_paths else "",
        "skill_file_paths": skill_paths,
        "file_path": profile_paths[0] if profile_paths else "",
        "file_paths": profile_paths,
        "model_file_path": skill_paths[0] if skill_paths else "",
        "model_file_paths": skill_paths,
        "department_assignment": assignment,
        "insert_mode": "append_non_destructive",
        "authority_level": str(authority_level).strip() or default_authority_level,
        "source_directory": str(directory),
        "scan_warning": scan_warning,
    }


def scan_member_directories(
    root_dir: str | Path,
    candidates: Iterable[str] | None = None,
    default_level: str = "专业成员",
    default_department: str = "产品与解决方案",
    recursive: bool = True,
    progress_callback: object | None = None,
) -> dict:
    """Scan a batch member directory.

    Expected layout:
        root/张三/01_基本信息.md
        root/张三/02_项目经历.md
        root/李四/01_基本信息.md

    Each first-level sub-directory is treated as one member. Files directly under
    root are skipped to avoid accidentally merging multiple people into one profile.
    """
    root = Path(root_dir)
    if not root.exists() or not root.is_dir():
        raise ValueError(f"批量成员目录不存在或不是文件夹：{root}")
    if callable(progress_callback):
        try:
            progress_callback("正在读取批量成员根目录...")
        except Exception:
            pass
    child_dirs: list[Path] = []
    root_supported_files: list[Path] = []
    hidden_or_excluded_count = 0
    root_entry_count = 0
    root_scan_truncated = False
    try:
        root_entries_iter = root.iterdir()
    except Exception as exc:
        raise ValueError(f"无法读取批量成员根目录：{exc}")
    for item in root_entries_iter:
        root_entry_count += 1
        if root_entry_count > BATCH_MEMBER_MAX_SCAN_ENTRIES_PER_MEMBER:
            root_scan_truncated = True
            break
        if callable(progress_callback) and root_entry_count % 50 == 0:
            try:
                progress_callback(f"正在扫描根目录条目：已检查 {root_entry_count} 项，发现 {len(child_dirs)} 个成员目录...")
            except Exception:
                pass
        try:
            if item.is_dir():
                if item.is_symlink() or _is_hidden_or_excluded_dir(item):
                    hidden_or_excluded_count += 1
                    continue
                child_dirs.append(item)
            elif item.is_file() and item.suffix.lower() in BATCH_MEMBER_FILE_TYPES:
                root_supported_files.append(item)
        except Exception:
            hidden_or_excluded_count += 1
            continue
    warnings: list[str] = []
    if root_scan_truncated:
        warnings.append(f"根目录扫描已截断：最多检查 {BATCH_MEMBER_MAX_SCAN_ENTRIES_PER_MEMBER} 个条目；请确认没有误选大目录。")
    if root_supported_files:
        warnings.append(f"根目录下有 {len(root_supported_files)} 个资料文件未导入；批量导入要求每个员工单独放在一个子文件夹。")
    if hidden_or_excluded_count:
        warnings.append(f"已跳过 {hidden_or_excluded_count} 个隐藏、软链接或排除目录。")
    child_dirs = sorted(child_dirs, key=lambda path: path.name.lower())
    if len(child_dirs) > BATCH_MEMBER_MAX_MEMBERS:
        warnings.append(f"发现 {len(child_dirs)} 个成员子文件夹，本次只导入前 {BATCH_MEMBER_MAX_MEMBERS} 个；请分批导入其余目录。")
        child_dirs = child_dirs[:BATCH_MEMBER_MAX_MEMBERS]
    if callable(progress_callback):
        try:
            progress_callback(f"发现 {len(child_dirs)} 个成员目录，开始解析...")
        except Exception:
            pass
    members: list[dict] = []
    skipped: list[str] = []
    for index, folder in enumerate(child_dirs, start=1):
        if callable(progress_callback):
            try:
                progress_callback(f"正在扫描成员目录 {index}/{len(child_dirs)}：{folder.name}")
            except Exception:
                pass
        try:
            member = build_member_payload_from_directory(
                folder,
                candidates=candidates,
                default_level=default_level,
                default_department=default_department,
                recursive=recursive,
                progress_callback=progress_callback,
            )
        except Exception as exc:
            skipped.append(f"{folder.name}：{exc}")
            continue
        if member.get("scan_warning"):
            warnings.append(f"{folder.name}：{member.get('scan_warning')}")
        members.append(member)
    if not child_dirs:
        warnings.append("未发现成员子文件夹；请使用“一个员工一个文件夹”的目录结构。")
    return {
        "root_dir": str(root),
        "members": members,
        "warnings": warnings,
        "skipped": skipped,
        "scanned_directory_count": len(child_dirs),
        "imported_member_count": len(members),
    }


def summarize_documents(file_paths: Iterable[str | Path], max_chars_per_file: int = 5000, max_total_chars: int = 18000) -> str:
    sections: list[str] = []
    total = 0
    for path in file_paths:
        content = extract_text_from_file(path, max_chars=max_chars_per_file).strip()
        if not content:
            continue
        header = f"### 文件：{Path(path).name}\n"
        block = header + content + "\n"
        if total + len(block) > max_total_chars:
            remaining = max_total_chars - total
            if remaining > 200:
                block = block[:remaining]
                sections.append(block)
            break
        sections.append(block)
        total += len(block)
    return "\n".join(sections).strip()


def search_latest_role_info(company_name: str, department: str, position: str, max_results: int = 5) -> list[dict]:
    role_query = " ".join([part for part in [company_name, department, position, "岗位职责 最新"] if part])
    query = urllib.parse.quote(role_query)
    url = f"https://duckduckgo.com/html/?q={query}"
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36"
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as response:
            html_text = response.read().decode("utf-8", errors="ignore")
    except Exception:
        return []

    pattern = re.compile(
        r'<a[^>]*class="result__a"[^>]*href="(?P<href>[^"]+)"[^>]*>(?P<title>.*?)</a>.*?(?:<a[^>]*class="result__snippet"[^>]*>|<div[^>]*class="result__snippet"[^>]*>)(?P<snippet>.*?)(?:</a>|</div>)',
        re.S,
    )

    results: list[dict] = []
    for match in pattern.finditer(html_text):
        title = clean_html(match.group("title"))
        snippet = clean_html(match.group("snippet"))
        href = html.unescape(match.group("href"))
        if title and href:
            results.append({"title": title, "snippet": snippet, "url": href})
        if len(results) >= max_results:
            break
    return results


def clean_html(raw_text: str) -> str:
    text = re.sub(r"<.*?>", " ", raw_text)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def format_search_results(results: list[dict]) -> str:
    if not results:
        return "未检索到可用的联网岗位信息，后续将改为基于上传资料和通用岗位知识生成。"
    lines = []
    for idx, item in enumerate(results, start=1):
        lines.append(f"{idx}. {item.get('title', '无标题')}")
        if item.get("snippet"):
            lines.append(f"   摘要：{item['snippet']}")
        if item.get("url"):
            lines.append(f"   链接：{item['url']}")
    return "\n".join(lines)


def format_organization_members_for_context(payload: dict) -> str:
    members = payload.get("organization_members") or []
    if not isinstance(members, list) or not members:
        return "暂无显式插入成员。"
    lines: list[str] = [
        "说明：以下成员是非覆盖式插入到组织架构中的组织插槽成员。它们不替换原老板 Skill、员工集合体 Skill 或公司行为准则。",
        "如果成员插入到老板位、董事会/创始层或高管层，回答时应作为插槽领导成员处理，而不是只作为外部顾问。",
    ]
    for idx, member in enumerate(members, start=1):
        if not isinstance(member, dict) or not member.get("enabled", True):
            continue
        assignment = member.get("department_assignment") or {}
        profile_paths = merge_file_path_values(member.get("profile_file_paths"), member.get("file_paths"), member.get("profile_file_path"), member.get("file_path"))
        skill_paths = merge_file_path_values(member.get("skill_file_paths"), member.get("model_file_paths"), member.get("skill_file_path"), member.get("model_file_path"))
        aliases = member.get("member_aliases") or []
        aliases_text = "、".join(aliases) if aliases else "无"
        lines.extend([
            f"{idx}. {member.get('member_name') or member.get('position') or '未命名成员'}",
            f"   - Skill 识别名：{member.get('skill_identity') or '未识别'}",
            f"   - 可匹配别名：{aliases_text}",
            f"   - 成员类型：{member.get('member_type') or '组织成员'}",
            f"   - 插入层级：{member.get('level') or '未设置'}",
            f"   - 插入部门：{member.get('department') or assignment.get('final_department') or '未设置'}",
            f"   - 插入岗位/位置：{member.get('position') or '未设置'}",
            f"   - 插入方式：{member.get('insert_mode') or 'append_non_destructive'}（非覆盖）",
            f"   - 权限级别：{member.get('authority_level') or 'advisory'}",
            f"   - 资料文件：{format_file_names(profile_paths)}",
            f"   - Skill/模型文件：{format_file_names(skill_paths)}",
            f"   - 部门归类方式：{assignment.get('mode') or '未设置'}；依据：{assignment.get('reason') or '未填写'}",
            f"   - 补充说明：{member.get('role_notes') or '无'}",
        ])
    return "\n".join(lines) if len(lines) > 2 else "暂无显式插入成员。"


def build_request_context(payload: dict, company_docs_summary: str, role_search_text: str) -> str:
    org = payload.get("organization", {})
    personal = payload.get("personal_distillation", {})
    organization_members_text = format_organization_members_for_context(payload)
    selected_departments = org.get("selected_departments") or DEFAULT_DEPARTMENTS
    custom_departments = org.get("custom_departments") or []
    company_name = payload.get("company_name") or payload.get("company_slug")
    context = f"""
# V2.1 任务上下文

## 项目版本
{payload.get('version', '2.1')}

## 公司标识
{payload.get('company_slug', '')}

## 公司名称
{company_name}

## 模型模式
{payload.get('model_type', 'cloud')}

## 组织架构设置
- 组织架构来源：{'上传组织架构资料优先' if org.get('use_uploaded_org_chart') else '界面勾选/默认部门'}
- 已识别组织架构资料：{', '.join(Path(path).name for path in org.get('org_chart_files', [])) if org.get('org_chart_files') else '无'}
- 标准部门：{', '.join(selected_departments) if selected_departments else '未手动勾选，优先参考上传组织架构资料'}
- 自定义部门：{', '.join(custom_departments) if custom_departments else '无'}
- 组织层级说明：{org.get('org_notes') or '未填写'}

## 组织人员蒸馏设置
- 是否启用：{'是' if personal.get('enabled') else '否'}
- 所在层级：{personal.get('level') or '未设置'}
- 所在部门：{personal.get('department') or '未设置'}
- 部门归类方式：{personal.get('department_assignment', {}).get('mode', 'manual')}
- 自动识别部门：{personal.get('department_assignment', {}).get('detected_department') or '无'}
- 自动识别置信度：{personal.get('department_assignment', {}).get('confidence_percent') or '无'}
- 最终采用部门：{personal.get('department_assignment', {}).get('final_department') or personal.get('department') or '未设置'}
- 识别/标注依据：{personal.get('department_assignment', {}).get('reason') or '未填写'}
- 所在岗位：{personal.get('position') or '未设置'}
- 岗位/人员补充说明：{personal.get('role_notes') or '未填写'}
- 组织人员资料文件：{format_file_names(merge_file_path_values(personal.get('file_paths'), personal.get('file_path')))}
- 组织人员蒸馏模型文件：{format_file_names(merge_file_path_values(personal.get('model_file_paths'), personal.get('model_file_path')))}
- 是否已有同组织其他人员蒸馏模型：{'是' if personal.get('has_peer_models', True) else '否'}

## 组织成员插入清单（非覆盖式）
{organization_members_text}

## 联网岗位信息
{role_search_text}

## 用户确认补充
{payload.get('confirmation_notes') or '暂无'}

## 公司资料摘要
{company_docs_summary or '暂无公司资料'}
""".strip()
    return context


def build_preview_prompt(context: str) -> str:
    return textwrap.dedent(
        f"""
        你是公司技能生成器 V2.1 的生成前确认助手。

        请基于以下上下文输出一份“生成确认建议”，帮助用户在正式生成前进一步优化配置。

        输出格式必须包含：
        1. 配置完整度判断
        2. 组织架构合理性检查
        3. 文件资料缺口
        4. 若没有同组织蒸馏模型，基于联网岗位信息总结该岗位的最新职责重点
        5. 给用户的3-5条优化建议
        6. 推荐最终生成方向（公司战略图景 / 组织人员发展预测 / 两者都生成）

        以下是上下文：
        {context}
        """
    ).strip()
